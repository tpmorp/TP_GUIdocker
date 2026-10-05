"""
TP_GUIdocker の画面文言・見た目・動作の設定を、1つの元データ (docs/ui_spec.json) から生成する。

  - src/ui_spec.h                     … プラグインが使う定義 (文言・数値)
  - docs/TP_GUIdocker_指示書.xlsx     … 利用者が修正指示を書き込む一覧表

使い方: py -3 tools/gen_ui_spec.py
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SPEC = ROOT / "docs" / "ui_spec.json"
HEADER = ROOT / "src" / "ui_spec.h"
WORKBOOK = ROOT / "docs" / "TP_GUIdocker_指示書.xlsx"

FONT_NAME = "Meiryo UI"


# ---------------------------------------------------------------------
#  C++ 定義の生成
# ---------------------------------------------------------------------

def cpp_wide_literal(s: str) -> str:
    out = []
    for ch in s:
        if ch == "\\":
            out.append("\\\\")
        elif ch == '"':
            out.append('\\"')
        elif ch == "\n":
            out.append("\\n")
        else:
            out.append(ch)
    return 'L"' + "".join(out) + '"'


def parse_color(value: str) -> int:
    m = re.fullmatch(r"#([0-9A-Fa-f]{6})", value.strip())
    if not m:
        raise ValueError(f"色は #RRGGBB で書くこと: {value}")
    rgb = int(m.group(1), 16)
    r, g, b = (rgb >> 16) & 0xFF, (rgb >> 8) & 0xFF, rgb & 0xFF
    return (b << 16) | (g << 8) | r  # COLORREF (0x00BBGGRR)


# 動作の選択肢 → 定義値
BEHAVIOR_CHOICES = {
    "DROP_STYLE": {"十字ボタン式": 0, "区画の領域式": 1},
    "DROP_OUTSIDE": {"浮いた窓にする": 1, "何もしない": 0},
    "EMPTY_AREA_CLOSE": {"自動で閉じて隣が広がる": 1, "自動で閉じる": 1, "空き区画として残す": 0},
}


def generate_header(spec: dict) -> str:
    lines = [
        "//----------------------------------------------------------------------------------",
        "//\tTP_GUIdocker 画面文言・見た目・動作の設定",
        "//\t※ このファイルは tools/gen_ui_spec.py が docs/ui_spec.json から生成する。直接編集しないこと",
        "//----------------------------------------------------------------------------------",
        "#pragma once",
        "#include <windows.h>",
        "",
        "namespace tpgd::ui {",
        "",
        "//---------------------------------------------------------------------",
        "//\t文言",
        "//---------------------------------------------------------------------",
        "",
    ]
    for t in spec["text"]:
        lines.append(f"// {t['where']} [{t['status']}]")
        lines.append(f"constexpr const wchar_t* {t['id']} = {cpp_wide_literal(t['value'])};")
    lines += ["", "//---------------------------------------------------------------------",
              "//\t見た目 (数値で指定するもの。本体設定に連動する項目は host.cpp で読む)",
              "//---------------------------------------------------------------------", ""]
    for d in spec["design"]:
        v = d["value"].strip()
        if d["unit"] == "色" and v.startswith("#"):
            lines.append(f"// {d['item']} ({v}) [{d['status']}]")
            lines.append(f"constexpr COLORREF {d['id']} = 0x{parse_color(v):06X};")
        elif re.fullmatch(r"-?\d+", v):
            lines.append(f"// {d['item']} ({d['unit']}) [{d['status']}]")
            lines.append(f"constexpr int {d['id']} = {v};")
        elif d["unit"] == "書体":
            face = "" if v == "メニューバーと同じ" else v
            lines.append(f"// {d['item']}: {v} (空: メニューバーと同じ書体) [{d['status']}]")
            lines.append(f"constexpr const wchar_t* {d['id']} = {cpp_wide_literal(face)};")
        elif v == "メニューバーと同じ":
            lines.append(f"// {d['item']}: メニューバーと同じ高さ (-1) [{d['status']}]")
            lines.append(f"constexpr int {d['id']} = -1;")
        else:
            lines.append(f"// {d['item']}: {v} (本体設定に連動)")
    lines += ["", "//---------------------------------------------------------------------",
              "//\t動作", "//---------------------------------------------------------------------", ""]
    for b in spec["behavior"]:
        v = b["value"].strip()
        if b["id"] in BEHAVIOR_CHOICES:
            choices = BEHAVIOR_CHOICES[b["id"]]
            if v not in choices:
                raise ValueError(f"{b['id']} の値 '{v}' は選択肢 {list(choices)} に無い")
            lines.append(f"// {b['item']}: {v} [{b['status']}]")
            lines.append(f"constexpr int {b['id']} = {choices[v]};")
        elif re.fullmatch(r"-?\d+", v):
            lines.append(f"// {b['item']} ({b['unit']}) [{b['status']}]")
            lines.append(f"constexpr int {b['id']} = {v};")
        else:
            lines.append(f"// {b['item']}: {v} [{b['status']}]")
    lines += ["", "} // namespace tpgd::ui", ""]
    return "\n".join(lines)


# ---------------------------------------------------------------------
#  指示書 (Excel) の生成
# ---------------------------------------------------------------------

THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="3A4A6B")
INPUT_FILL = PatternFill("solid", fgColor="FFF59D")      # 書き込み欄 (黄)
DONE_FILL = PatternFill("solid", fgColor="E8F5E9")       # 実装済み
PLAN_FILL = PatternFill("solid", fgColor="FFF3E0")       # 段階2で実装
WRAP = Alignment(wrap_text=True, vertical="top")


def font(bold=False, color="000000", size=10):
    return Font(name=FONT_NAME, bold=bold, color=color, size=size)


def write_table(ws, headers, widths, rows, input_col, status_col):
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 22
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = font()
            cell.alignment = WRAP
            cell.border = BORDER
        ws.cell(row=r, column=input_col).fill = INPUT_FILL
        status = ws.cell(row=r, column=status_col)
        status.fill = DONE_FILL if status.value == "実装済み" else PLAN_FILL
        # 改行の多い文言は行を高くする
        longest = max(str(v).count("\n") + 1 + len(str(v)) // 40 for v in row if v is not None)
        ws.row_dimensions[r].height = max(18, min(150, 15 * longest))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


INPUT_HEADER = "修正指示"


def read_existing_instructions(path: Path) -> dict:
    """既存の指示書から、シートごと・ID ごとの「修正指示」を読み取る (作り直しても消さないため)"""
    found = {}
    if not path.exists():
        return found
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True)
    except Exception as e:  # noqa: BLE001
        print(f"注意: 既存の指示書を読めない ({e})")
        return found
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows or INPUT_HEADER not in rows[0]:
            continue
        col = rows[0].index(INPUT_HEADER)
        for row in rows[1:]:
            if row and row[0] and col < len(row) and row[col] not in (None, ""):
                found.setdefault(ws.title, {})[row[0]] = row[col]
    return found


def drop_applied(kept: dict, spec: dict) -> dict:
    """反映の記録に「反映済み」として載った指示は、修正指示欄から外す"""
    done = {(h["sheet"], h["id"], str(h["instruction"])) for h in spec.get("history", []) if h.get("result") == "反映済み"}
    out = {}
    for sheet, items in kept.items():
        for item_id, text in items.items():
            if (sheet, item_id, str(text)) in done:
                continue
            out.setdefault(sheet, {})[item_id] = text
    return out


def generate_workbook(spec: dict, kept: dict) -> Workbook:
    wb = Workbook()

    # 使い方
    ws = wb.active
    ws.title = "使い方"
    ws.column_dimensions["A"].width = 100
    guide = [
        ("TP_GUIdocker 指示書", font(bold=True, size=14)),
        ("", None),
        ("この表は、TP_GUIdocker の画面に出る文言・見た目・動作を一覧にしたものです。", font()),
        ("変えたいものがあれば、黄色の「修正指示」欄に書き込んでください。それ以外の欄は変更不要です。", font(bold=True)),
        ("", None),
        ("■ シート", font(bold=True)),
        ("　文言 … メニュー・案内・確認画面などに出る文字", font()),
        ("　見た目 … 大きさ・余白・色など", font()),
        ("　動作 … 長押しの時間・落とし先の見せ方など", font()),
        ("", None),
        ("■ 「状態」欄", font(bold=True)),
        ("　実装済み (緑) … 現在のプラグインで使われている", font()),
        ("　段階2で実装 (橙) … これから作る部分。作る前に指示いただければ最初から反映する", font()),
        ("", None),
        ("■ 書き方の例", font(bold=True)),
        ("　文言 MENU_RESET の「修正指示」欄に → 「配置をリセット」", font()),
        ("　見た目 COMPASS_BUTTON の「修正指示」欄に → 「64 にして、もう少し大きく」", font()),
        ("　色は #RRGGBB (例: #FF8800) か、「もう少し明るい青」のような言葉でも構いません。", font()),
        ("　文言の中の %s・%d・%lu は、実際の名前や数に置き換わる部分です。消さずに残してください。", font()),
        ("", None),
        ("■ 書き込んだ後", font(bold=True)),
        ("　保存して Excel を閉じ、書き込んだことを知らせてください。内容を反映したプラグインを作り直します。", font()),
        ("　反映した指示は「修正指示」欄から消え、「反映の記録」シートに残ります。相談中の指示は欄に残ります。", font()),
        ("　(元データは docs/ui_spec.json。この表はそこから自動で作られます)", font(color="707070")),
    ]
    for i, (text, f) in enumerate(guide, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if f:
            c.font = f
    ws.cell(row=4, column=1).fill = INPUT_FILL

    def kept_of(sheet, item_id):
        return kept.get(sheet, {}).get(item_id)

    # 文言
    ws = wb.create_sheet("文言")
    write_table(ws,
                ["ID", "表示される場所", "現在の文言", INPUT_HEADER, "状態"],
                [24, 46, 60, 46, 14],
                [[t["id"], t["where"], t["value"], kept_of("文言", t["id"]), t["status"]] for t in spec["text"]],
                input_col=4, status_col=5)

    # 見た目
    ws = wb.create_sheet("見た目")
    write_table(ws,
                ["ID", "項目", "現在の値", "単位", INPUT_HEADER, "状態", "備考"],
                [22, 36, 34, 22, 40, 14, 34],
                [[d["id"], d["item"], d["value"], d["unit"], kept_of("見た目", d["id"]), d["status"], d["note"]] for d in spec["design"]],
                input_col=5, status_col=6)

    # 動作
    ws = wb.create_sheet("動作")
    write_table(ws,
                ["ID", "項目", "現在の値", "選べる範囲", "単位", INPUT_HEADER, "状態", "備考"],
                [20, 36, 26, 30, 12, 40, 14, 34],
                [[b["id"], b["item"], b["value"], b["choices"], b["unit"], kept_of("動作", b["id"]), b["status"], b["note"]] for b in spec["behavior"]],
                input_col=6, status_col=7)

    # 反映の記録 (新しいものが上)
    ws = wb.create_sheet("反映の記録")
    headers = ["日付", "シート", "ID", "いただいた指示", "対応", "結果"]
    widths = [12, 10, 22, 50, 50, 12]
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    for r, h in enumerate(reversed(spec.get("history", [])), start=2):
        for c, v in enumerate([h["date"], h["sheet"], h["id"], h["instruction"], h["action"], h["result"]], start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = font()
            cell.alignment = WRAP
            cell.border = BORDER
        ws.cell(row=r, column=6).fill = DONE_FILL if h["result"] == "反映済み" else PLAN_FILL
    ws.freeze_panes = "A2"

    return wb


def main():
    spec = json.loads(SPEC.read_text(encoding="utf-8"))
    ids = [x["id"] for k in ("text", "design", "behavior") for x in spec[k]]
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        sys.exit(f"ID が重複している: {sorted(dup)}")

    HEADER.write_text(generate_header(spec), encoding="utf-8", newline="\r\n")
    print(f"生成: {HEADER}")

    kept = drop_applied(read_existing_instructions(WORKBOOK), spec)
    n = sum(len(v) for v in kept.values())
    if n:
        print(f"既存の指示書の「修正指示」{n} 件を引き継ぐ")
    try:
        generate_workbook(spec, kept).save(WORKBOOK)
        print(f"生成: {WORKBOOK}")
    except PermissionError:
        print(f"注意: 指示書が開かれているため更新を見送った ({WORKBOOK.name})。閉じてから作り直すこと")


if __name__ == "__main__":
    main()
